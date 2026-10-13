"""DNP3 outstation points list: model, validation and spreadsheet import/export.

The points list is the single source of truth for the outstation database and
for what each diagram element can bind to. It is plain Python (no Qt) so it can
be used headless, e.g. for tests or a console-only outstation.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field, asdict
from pathlib import Path

# Point types, using DNP3 terminology from the outstation's point of view.
BI = "BI"  # Binary Input          - Group 1 static / Group 2 events
BO = "BO"  # Binary Output         - Group 10 status / Group 12 CROB control
AI = "AI"  # Analog Input          - Group 30 static / Group 32 events
AO = "AO"  # Analog Output         - Group 40 status / Group 41 command
POINT_TYPES = (BI, BO, AI, AO)

# How a binary output is operated. Each model accepts the CROB codes a master
# would normally send for it, and drives the HMI's button captions.
CONTROL_MODELS = {
    "TripClose": ("Trip", "Close"),
    "OpenClose": ("Open", "Close"),
    "OnOff": ("Off", "On"),
}

COLUMNS = [
    "name", "type", "index", "description", "units", "event_class",
    "control_model", "feedback", "value", "min", "max",
]
HEADERS = {
    "name": "Name", "type": "Type", "index": "Index", "description": "Description",
    "units": "Units", "event_class": "Class", "control_model": "Control Model",
    "feedback": "Feedback Point", "value": "Value", "min": "Min", "max": "Max",
}


@dataclass
class Point:
    name: str
    type: str
    index: int
    description: str = ""
    units: str = ""
    event_class: int = 1
    control_model: str = ""   # BO only, one of CONTROL_MODELS
    feedback: str = ""        # BO -> BI, AO -> AI: point updated when operated
    value: float = 0.0
    min: float | None = None  # AO limits; commands outside are rejected
    max: float | None = None

    @property
    def is_binary(self) -> bool:
        return self.type in (BI, BO)

    @property
    def state(self) -> bool:
        return bool(self.value)

    def captions(self) -> tuple[str, str]:
        """(off, on) captions for a binary point."""
        return CONTROL_MODELS.get(self.control_model or "OpenClose", ("Off", "On"))

    def state_text(self, model: str | None = None) -> str:
        """State wording (not command wording): CLOSED/OPEN or ON/OFF."""
        model = model if model is not None else self.control_model
        if model in ("TripClose", "OpenClose"):
            return "CLOSED" if self.state else "OPEN"
        return "ON" if self.state else "OFF"

    def display_value(self) -> str:
        if self.is_binary:
            return self.state_text()
        text = f"{self.value:.2f}".rstrip("0").rstrip(".")
        return f"{text} {self.units}".strip()

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict) -> "Point":
        return cls(**{k: d[k] for k in COLUMNS if k in d})


def _blank(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def coerce_point(raw: dict) -> Point:
    """Build a Point from loosely typed spreadsheet cells. Raises ValueError."""
    def text(key):
        v = raw.get(key)
        return "" if _blank(v) else str(v).strip()

    def number(key, default=None):
        v = raw.get(key)
        if _blank(v):
            return default
        if isinstance(v, str) and v.strip().lower() in ("true", "on", "closed", "close"):
            return 1.0
        if isinstance(v, str) and v.strip().lower() in ("false", "off", "open", "trip"):
            return 0.0
        return float(v)

    name = text("name")
    if not name:
        raise ValueError("missing name")
    ptype = text("type").upper()
    if ptype not in POINT_TYPES:
        raise ValueError(f"{name}: type must be one of {', '.join(POINT_TYPES)}, got '{ptype}'")
    index = number("index")
    if index is None or index < 0 or int(index) != index:
        raise ValueError(f"{name}: index must be a whole number >= 0")
    model = text("control_model")
    if ptype == BO:
        model = model or "OpenClose"
        match = {m.lower(): m for m in CONTROL_MODELS}.get(model.lower().replace("/", ""))
        if not match:
            raise ValueError(f"{name}: control model must be one of {', '.join(CONTROL_MODELS)}")
        model = match
    else:
        model = ""
    event_class = int(number("event_class", 1))
    if event_class not in (0, 1, 2, 3):
        raise ValueError(f"{name}: class must be 0, 1, 2 or 3")
    return Point(
        name=name, type=ptype, index=int(index), description=text("description"),
        units=text("units"), event_class=event_class, control_model=model,
        feedback=text("feedback"), value=number("value", 0.0) or 0.0,
        min=number("min"), max=number("max"),
    )


@dataclass
class PointList:
    points: list[Point] = field(default_factory=list)

    def __iter__(self):
        return iter(self.points)

    def __len__(self):
        return len(self.points)

    def get(self, name: str) -> Point | None:
        for p in self.points:
            if p.name == name:
                return p
        return None

    def by_index(self, ptype: str, index: int) -> Point | None:
        for p in self.points:
            if p.type == ptype and p.index == index:
                return p
        return None

    def state_text(self, p: Point) -> str:
        """A BI used as feedback takes its wording from the BO that drives it."""
        if p.type == BI:
            for q in self.points:
                if q.type == BO and q.feedback == p.name:
                    return p.state_text(q.control_model)
        return p.state_text()

    def names(self, *types: str) -> list[str]:
        return [p.name for p in self.points if not types or p.type in types]

    def next_index(self, ptype: str) -> int:
        used = [p.index for p in self.points if p.type == ptype]
        return max(used) + 1 if used else 0

    def unique_name(self, base: str) -> str:
        n, name = 1, base
        while self.get(name):
            n += 1
            name = f"{base}_{n}"
        return name

    def validate(self) -> list[str]:
        """Return human-readable problems; an empty list means valid."""
        problems = []
        seen_names, seen_idx = set(), set()
        for p in self.points:
            if p.name in seen_names:
                problems.append(f"Duplicate point name '{p.name}'")
            seen_names.add(p.name)
            if (p.type, p.index) in seen_idx:
                problems.append(f"Duplicate {p.type} index {p.index} ('{p.name}')")
            seen_idx.add((p.type, p.index))
            if p.feedback:
                fb = self.get(p.feedback)
                want = {BO: BI, AO: AI}.get(p.type)
                if want is None:
                    problems.append(f"'{p.name}': only BO and AO points can have a feedback point")
                elif fb is None:
                    problems.append(f"'{p.name}': feedback point '{p.feedback}' does not exist")
                elif fb.type != want:
                    problems.append(f"'{p.name}': feedback point must be a {want}, '{fb.name}' is {fb.type}")
        return problems

    def sizes(self) -> dict[str, int]:
        """DNP3 database size per type (highest index + 1, gaps allowed)."""
        return {t: max([p.index for p in self.points if p.type == t], default=-1) + 1
                for t in POINT_TYPES}

    # --- serialisation -------------------------------------------------
    def to_list(self) -> list[dict]:
        return [p.to_dict() for p in self.points]

    @classmethod
    def from_list(cls, rows: list[dict]) -> "PointList":
        return cls([Point.from_dict(r) for r in rows])

    # --- spreadsheet import/export --------------------------------------
    def export_file(self, path: str | Path) -> None:
        path = Path(path)
        rows = [[_cell(getattr(p, c)) for c in COLUMNS] for p in self.points]
        if path.suffix.lower() == ".csv":
            with open(path, "w", newline="") as f:
                w = csv.writer(f)
                w.writerow([HEADERS[c] for c in COLUMNS])
                w.writerows(rows)
            return
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation

        wb = Workbook()
        ws = wb.active
        ws.title = "Points"
        ws.append([HEADERS[c] for c in COLUMNS])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F4858")
        for r in rows:
            ws.append(r)
        widths = {"name": 18, "description": 32, "control_model": 14, "feedback": 18}
        for i, c in enumerate(COLUMNS):
            ws.column_dimensions[chr(65 + i)].width = widths.get(c, 10)
        ws.freeze_panes = "A2"
        # Drop-downs so the sheet stays importable when edited by hand.
        last = max(len(rows) + 1, 500)
        for col, options in (("type", POINT_TYPES), ("control_model", tuple(CONTROL_MODELS))):
            letter = chr(65 + COLUMNS.index(col))
            dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
            dv.add(f"{letter}2:{letter}{last}")
            ws.add_data_validation(dv)
        wb.save(path)

    @classmethod
    def import_file(cls, path: str | Path) -> tuple["PointList", list[str]]:
        """Read a points sheet. Returns (points, problems). Bad rows are skipped."""
        path = Path(path)
        if path.suffix.lower() == ".csv":
            with open(path, newline="", encoding="utf-8-sig") as f:
                table = list(csv.reader(f))
        else:
            from openpyxl import load_workbook
            wb = load_workbook(path, data_only=True, read_only=True)
            ws = wb["Points"] if "Points" in wb.sheetnames else wb.active
            table = [list(r) for r in ws.iter_rows(values_only=True)]
        if not table:
            return cls(), ["The file is empty"]

        # Match headers by display name or field name, ignoring case and spaces.
        lookup = {}
        for c in COLUMNS:
            lookup[c.replace("_", "")] = c
            lookup[HEADERS[c].lower().replace(" ", "")] = c
        lookup.update({"class": "event_class", "feedbackpoint": "feedback", "model": "control_model"})
        header = [lookup.get(str(h or "").strip().lower().replace(" ", "").replace("_", "")) for h in table[0]]
        missing = [HEADERS[c] for c in ("name", "type", "index") if c not in header]
        if missing:
            return cls(), [f"Missing required column(s): {', '.join(missing)}"]

        points, problems = cls(), []
        for rownum, row in enumerate(table[1:], start=2):
            if all(_blank(v) for v in row):
                continue
            raw = {key: row[i] for i, key in enumerate(header) if key and i < len(row)}
            try:
                points.points.append(coerce_point(raw))
            except (ValueError, TypeError) as e:
                problems.append(f"Row {rownum}: {e}")
        problems += points.validate()
        return points, problems


def _cell(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v
